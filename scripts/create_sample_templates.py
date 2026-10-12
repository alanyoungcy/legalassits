#!/usr/bin/env python3
"""Generate corporate sample Word templates with yellow fillable fields."""

from __future__ import annotations

import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
TEMPLATES = ROOT / "templates"


def add_yellow(paragraph, text: str) -> None:
    run = paragraph.add_run(text)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    run.font.size = Pt(11)


def add_plain(paragraph, text: str, *, bold: bool = False) -> None:
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(11)


def write_meta(docx_path: Path, meta: dict) -> None:
    meta_path = docx_path.with_suffix(".meta.json")
    meta_path.write_text(json.dumps(meta, indent=2), encoding="utf-8")


def build_nda() -> Path:
    doc = Document()
    title = doc.add_heading("", level=1)
    add_plain(title, "Mutual Non-Disclosure Agreement", bold=True)

    p = doc.add_paragraph()
    add_plain(p, "This Mutual Non-Disclosure Agreement (the “Agreement”) is entered into as of ")
    add_yellow(p, "[Effective Date]")
    add_plain(p, " (the “Effective Date”), by and between:")

    p = doc.add_paragraph()
    add_yellow(p, "[Disclosing Party]")
    add_plain(p, ", a company organized under the laws of ")
    add_yellow(p, "[Governing Law]")
    add_plain(p, " with its principal place of business at ")
    add_yellow(p, "[Disclosing Party Address]")
    add_plain(p, " (“Disclosing Party”); and")

    p = doc.add_paragraph()
    add_yellow(p, "[Receiving Party]")
    add_plain(p, ", with its principal place of business at ")
    add_yellow(p, "[Receiving Party Address]")
    add_plain(p, " (“Receiving Party”).")

    doc.add_heading("1. Purpose", level=2)
    p = doc.add_paragraph()
    add_plain(p, "The parties wish to explore ")
    add_yellow(p, "[Purpose of Disclosure]")
    add_plain(p, " and may share confidential information for that purpose.")

    doc.add_heading("2. Term", level=2)
    p = doc.add_paragraph()
    add_plain(p, "This Agreement remains in effect for ")
    add_yellow(p, "[Confidentiality Term]")
    add_plain(p, " from the Effective Date, unless earlier terminated.")

    doc.add_heading("3. Capco entity", level=2)
    p = doc.add_paragraph()
    add_plain(p, "Capco contracting entity: ")
    add_yellow(p, "[Capco Legal Entity]")
    add_plain(p, ".")

    path = TEMPLATES / "corporate" / "nda_mutual.docx"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    write_meta(
        path,
        {
            "id": "nda-mutual",
            "name": "Mutual NDA",
            "agreement_type": "NDA",
            "jurisdiction": "Hong Kong / UK",
            "entity_scope": "Capco group",
            "description": (
                "Corporate mutual NDA. Yellow-highlighted placeholders are fillable fields "
                "indexed by Capco Contract Assists."
            ),
        },
    )
    return path


def build_msa() -> Path:
    doc = Document()
    title = doc.add_heading("", level=1)
    add_plain(title, "Master Services Agreement", bold=True)

    p = doc.add_paragraph()
    add_plain(p, "Agreement date: ")
    add_yellow(p, "{{effective_date}}")

    p = doc.add_paragraph()
    add_plain(p, "Service provider: ")
    add_yellow(p, "{{capco_legal_entity}}")
    add_plain(p, " (“Provider”).")

    p = doc.add_paragraph()
    add_plain(p, "Client: ")
    add_yellow(p, "{{client_name}}")
    add_plain(p, ", registered at ")
    add_yellow(p, "{{client_address}}")
    add_plain(p, " (“Client”).")

    doc.add_heading("1. Services", level=2)
    p = doc.add_paragraph()
    add_plain(p, "Provider will deliver ")
    add_yellow(p, "{{services_description}}")
    add_plain(p, " as further described in one or more Statements of Work.")

    doc.add_heading("2. Fees", level=2)
    p = doc.add_paragraph()
    add_plain(p, "Fees are payable in ")
    add_yellow(p, "{{currency}}")
    add_plain(p, ". Payment terms: ")
    add_yellow(p, "{{payment_terms}}")
    add_plain(p, ".")

    doc.add_heading("3. Governing law", level=2)
    p = doc.add_paragraph()
    add_plain(p, "This Agreement is governed by the laws of ")
    add_yellow(p, "{{governing_law}}")
    add_plain(p, ".")

    path = TEMPLATES / "corporate" / "msa_master_services.docx"
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    write_meta(
        path,
        {
            "id": "msa-master-services",
            "name": "Master Services Agreement",
            "agreement_type": "MSA",
            "jurisdiction": "United Kingdom",
            "entity_scope": "Capco UK",
            "description": "Corporate MSA using {{field_key}} yellow placeholders.",
        },
    )
    return path


def main() -> int:
    TEMPLATES.mkdir(parents=True, exist_ok=True)
    created = [build_nda(), build_msa()]
    readme = TEMPLATES / "README.md"
    readme.write_text(
        """# Corporate contract templates

Drop Capco / Wipro Word (`.docx`) templates here. Nested folders are supported.

## Yellow fields

Any text highlighted **yellow** in Word becomes a fillable field:

| Placeholder style in Word | Indexed field key |
| --- | --- |
| `[Party Name]` | `party_name` |
| `{{effective_date}}` | `effective_date` |
| `[[CONFIDENTIALITY_TERM]]` | `confidentiality_term` |
| Free text highlighted yellow | slug of that text |

Optional sidecar metadata: place `my-template.meta.json` next to `my-template.docx`.

```json
{
  "id": "nda-mutual",
  "name": "Mutual NDA",
  "agreement_type": "NDA",
  "jurisdiction": "Hong Kong",
  "entity_scope": "Capco group",
  "description": "Standard mutual NDA"
}
```

After adding files, call `POST /api/v1/templates/reindex` or restart the API.
""",
        encoding="utf-8",
    )
    print("Created:")
    for path in created:
        print(f"  - {path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
