"""Skill-guided contract review runner."""

from __future__ import annotations

import json
import re
import uuid
from datetime import datetime, timezone
from pathlib import Path

from docx import Document

from app.config import Settings, get_settings
from app.services.catalog import catalog
from app.services.contracts import contract_store
from app.services.skills import skill_registry


SEVERITY_HINTS = (
    ("red", "RED"),
    ("yellow", "YELLOW"),
    ("green", "GREEN"),
    ("escalat", "ESCALATE"),
    ("auto-yellow", "YELLOW"),
    ("provisional", "PROVISIONAL"),
)


class ReviewService:
    def __init__(self, settings: Settings | None = None) -> None:
        self.settings = settings or get_settings()
        self.index_path = self.settings.data_dir / "reviews_index.json"
        self._items: dict[str, dict] = {}
        self._load()

    def _load(self) -> None:
        if not self.index_path.exists():
            self._items = {}
            return
        self._items = json.loads(self.index_path.read_text(encoding="utf-8"))

    def _save(self) -> None:
        self.index_path.parent.mkdir(parents=True, exist_ok=True)
        self.index_path.write_text(json.dumps(self._items, indent=2), encoding="utf-8")

    def list_reviews(self, contract_id: str | None = None) -> list[dict]:
        items = list(self._items.values())
        if contract_id:
            items = [item for item in items if item.get("contract_id") == contract_id]
        return sorted(items, key=lambda row: row["created_at"], reverse=True)

    def get(self, review_id: str) -> dict | None:
        return self._items.get(review_id)

    def run(
        self,
        *,
        contract_id: str,
        skill_id: str | None = None,
        side: str | None = None,
    ) -> dict:
        contract = contract_store.get(contract_id)
        if not contract:
            raise KeyError(f"Unknown contract: {contract_id}")
        template = catalog.get(contract.template_id)
        agreement_type = template.agreement_type if template else None
        skill = skill_registry.route_review_skill(
            agreement_type=agreement_type,
            title=contract.title,
            explicit_skill_id=skill_id,
        )
        playbook = self._load_playbook(skill["plugin"])
        document_text = self._extract_docx_text(Path(contract.download_path) if contract.download_path else None)
        checklist = self._build_checklist(skill, playbook)
        routing = {
            "agreement_type": agreement_type,
            "selected_skill": skill["id"],
            "plugin": skill["plugin"],
            "side": side or "purchasing",
            "confirm_routing": True,
            "note": (
                "Draft for attorney review — not legal advice. "
                "Outputs follow anthropics/claude-for-legal skill workflows adapted for Capco."
            ),
        }
        findings = self._heuristic_findings(
            document_text=document_text,
            agreement_type=agreement_type,
            skill_id=skill["id"],
            field_values=contract.field_values,
        )
        memo = self._render_memo(
            contract_title=contract.title,
            routing=routing,
            skill=skill,
            checklist=checklist,
            findings=findings,
            playbook_excerpt=playbook[:1800],
        )
        review_id = f"rev_{uuid.uuid4().hex[:12]}"
        now = datetime.now(timezone.utc).isoformat()
        row = {
            "id": review_id,
            "contract_id": contract.id,
            "contract_title": contract.title,
            "skill_id": skill["id"],
            "skill_name": skill["name"],
            "plugin": skill["plugin"],
            "status": "draft_for_attorney_review",
            "triage": findings.get("triage"),
            "created_at": now,
            "updated_at": now,
            "routing": routing,
            "checklist": checklist,
            "findings": findings.get("items", []),
            "memo": memo,
            "disclaimer": (
                "Every output is a draft for attorney review — not legal advice, "
                "not a legal conclusion, and not a substitute for a lawyer. "
                "Source: anthropics/claude-for-legal skill pack."
            ),
        }
        self._items[review_id] = row
        self._save()
        return row

    def _load_playbook(self, plugin: str) -> str:
        path = self.settings.skills_playbooks_dir / f"{plugin}.CLAUDE.md"
        if path.exists():
            return path.read_text(encoding="utf-8")
        return ""

    @staticmethod
    def _extract_docx_text(path: Path | None) -> str:
        if not path or not path.exists():
            return ""
        document = Document(str(path))
        parts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    text = cell.text.strip()
                    if text:
                        parts.append(text)
        return "\n".join(parts)

    def _build_checklist(self, skill: dict, playbook: str) -> list[dict]:
        items: list[dict] = []
        for heading in skill.get("headings") or []:
            severity = "INFO"
            lowered = heading.lower()
            for needle, label in SEVERITY_HINTS:
                if needle in lowered:
                    severity = label
                    break
            items.append(
                {
                    "section": heading,
                    "severity": severity,
                    "source": "skill",
                    "instruction": f"Apply skill section “{heading}” against the agreement and Capco playbook.",
                }
            )
        # Pull NDA triage bullets from Capco playbook when relevant
        if "nda" in skill["id"] or skill["id"].endswith("/review"):
            for line in playbook.splitlines():
                if line.strip().startswith("- "):
                    items.append(
                        {
                            "section": "Capco playbook position",
                            "severity": "YELLOW" if "YELLOW" in line.upper() else (
                                "RED" if "RED" in line.upper() else "INFO"
                            ),
                            "source": "playbook",
                            "instruction": line.strip()[2:],
                        }
                    )
        # De-dupe while preserving order
        seen: set[str] = set()
        unique = []
        for item in items:
            key = f"{item['section']}|{item['instruction']}"
            if key in seen:
                continue
            seen.add(key)
            unique.append(item)
        return unique[:60]

    def _heuristic_findings(
        self,
        *,
        document_text: str,
        agreement_type: str | None,
        skill_id: str,
        field_values: dict[str, str],
    ) -> dict:
        text = document_text.lower()
        items: list[dict] = []
        triage = "YELLOW"

        def add(severity: str, title: str, detail: str, citation: str | None = None) -> None:
            items.append(
                {
                    "severity": severity,
                    "title": title,
                    "detail": detail,
                    "citation": citation,
                    "recommendation": "Attorney review required before circulation or signature.",
                }
            )

        if "nda" in skill_id or (agreement_type or "").upper() == "NDA":
            non_nda = [
                ("non-solicit", "Non-solicit covenant found in NDA-labelled document"),
                ("non-compete", "Non-compete covenant found in NDA-labelled document"),
                ("exclusiv", "Exclusivity language found in NDA-labelled document"),
                ("intellectual property", "IP assignment language found in NDA-labelled document"),
                ("right of first", "ROFR-style language found in NDA-labelled document"),
            ]
            for needle, title in non_nda:
                if needle in text:
                    add("YELLOW", title, "Scope creep beyond confidentiality — auto-YELLOW per NDA skill.")
                    triage = "YELLOW"
            if "residual" in text:
                add("YELLOW", "Residuals clause present", "Check against Capco NDA triage positions.")
            if re.search(r"\b(perpetual|indefinite)\b", text):
                add("YELLOW", "Perpetual / indefinite confidentiality", "Confirm survival period vs playbook.")
            if not items:
                add(
                    "GREEN",
                    "No automated scope-creep flags",
                    "Heuristic scan found no non-NDA covenants. Still confirm mutuality, term, carveouts, and governing law against playbook before GREEN.",
                )
                triage = "PROVISIONAL"

        elif "employment" in skill_id:
            if "non-compete" in text:
                add("YELLOW", "Restrictive covenant: non-compete", "Run jurisdiction enforceability check.")
            if "non-solicit" in text:
                add("YELLOW", "Restrictive covenant: non-solicit", "Confirm scope, duration, and consideration.")
            if not items:
                add(
                    "INFO",
                    "Employment review package prepared",
                    "Complete jurisdiction, classification, and covenant checks per hiring/termination skill.",
                )
                triage = "YELLOW"
        else:
            if "unlimited liability" in text or "liability is unlimited" in text:
                add("RED", "Unlimited liability language", "Escalate per commercial playbook.")
                triage = "RED"
            if "auto-renew" in text or "automatically renew" in text:
                add("YELLOW", "Auto-renewal present", "Capture cancel-by date in renewal register.")
            if not items:
                add(
                    "INFO",
                    "Vendor / MSA review package prepared",
                    "Run full vendor-agreement-review checklist against Capco purchasing/sales playbook.",
                )
                triage = "YELLOW"

        # Missing filled fields can block confident review
        missing = [key for key, value in field_values.items() if not value]
        if missing:
            add(
                "YELLOW",
                "Incomplete template fields",
                f"Missing values: {', '.join(missing[:12])}",
            )
            if triage == "GREEN":
                triage = "YELLOW"

        return {"triage": triage, "items": items}

    @staticmethod
    def _render_memo(
        *,
        contract_title: str,
        routing: dict,
        skill: dict,
        checklist: list[dict],
        findings: dict,
        playbook_excerpt: str,
    ) -> str:
        lines = [
            f"# Review memo — {contract_title}",
            "",
            f"**Triage:** {findings.get('triage', 'YELLOW')}",
            f"**Skill:** `{skill['id']}` ({skill['name']})",
            f"**Plugin:** {skill['plugin']}",
            f"**Side assumption:** {routing.get('side')}",
            "",
            "> Draft for attorney review — not legal advice.",
            "",
            "## Routing",
            f"- Agreement type: {routing.get('agreement_type') or 'unspecified'}",
            f"- Selected skill: {routing.get('selected_skill')}",
            "",
            "## Findings",
        ]
        for item in findings.get("items", []):
            lines.append(f"- **{item['severity']} — {item['title']}:** {item['detail']}")
        lines.extend(["", "## Checklist (from skill + Capco playbook)", ""])
        for item in checklist[:25]:
            lines.append(f"- [{item['severity']}] {item['section']}: {item['instruction']}")
        if playbook_excerpt:
            lines.extend(["", "## Capco playbook excerpt", "", "```", playbook_excerpt, "```"])
        lines.extend(
            [
                "",
                "## Skill description",
                "",
                skill.get("description") or "",
                "",
                "## Next steps",
                "- Attorney verifies citations and playbook positions",
                "- Confirm or correct routing if confirm_routing is enabled",
                "- Escalate via commercial-legal/escalation-flagger when authority is exceeded",
            ]
        )
        return "\n".join(lines)


review_service = ReviewService()
