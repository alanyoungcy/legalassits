"""Extract and fill yellow-highlighted fields in corporate Word templates.

Corporate authors mark placeholders in Word with yellow highlight. This module:
1. Walks paragraphs and table cells
2. Collects contiguous yellow runs as a single field occurrence
3. Derives a stable field key from the placeholder text
4. Fills values and clears the highlight on export
"""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph


YELLOW = WD_COLOR_INDEX.YELLOW

# Explicit key forms: {{party_name}}, [[PARTY_NAME]], [PARTY NAME]
KEY_PATTERNS = (
    re.compile(r"^\{\{\s*([a-zA-Z0-9_.-]+)\s*\}\}$"),
    re.compile(r"^\[\[\s*([a-zA-Z0-9_.-]+)\s*\]\]$"),
    re.compile(r"^\[\s*([A-Za-z0-9_ .-]+)\s*\]$"),
)


def slugify_key(text: str) -> str:
    cleaned = text.strip().lower()
    cleaned = re.sub(r"[\{\}\[\]<>]", "", cleaned)
    cleaned = re.sub(r"[^a-z0-9]+", "_", cleaned)
    cleaned = cleaned.strip("_")
    return cleaned or "field"


def derive_field_key(placeholder: str) -> tuple[str, str]:
    """Return (key, human_label) from yellow placeholder text."""
    raw = " ".join(placeholder.split()).strip()
    for pattern in KEY_PATTERNS:
        match = pattern.match(raw)
        if match:
            key = slugify_key(match.group(1))
            label = match.group(1).replace("_", " ").strip().title()
            return key, label
    key = slugify_key(raw)
    label = raw if raw else key.replace("_", " ").title()
    return key, label


def _is_yellow_run(run) -> bool:
    try:
        if run.font.highlight_color == YELLOW:
            return True
    except Exception:
        pass
    # Fallback: inspect raw XML highlight value
    rpr = run._element.find(qn("w:rPr"))
    if rpr is None:
        return False
    highlight = rpr.find(qn("w:highlight"))
    if highlight is None:
        return False
    return (highlight.get(qn("w:val")) or "").lower() == "yellow"


def _iter_paragraphs(document: Document) -> Iterable[Paragraph]:
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        yield from _iter_table_paragraphs(table)


def _iter_table_paragraphs(table: Table) -> Iterable[Paragraph]:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                yield paragraph
            for nested in cell.tables:
                yield from _iter_table_paragraphs(nested)


@dataclass
class FieldOccurrence:
    key: str
    label: str
    placeholder: str
    paragraph_index: int
    run_start: int
    run_end: int  # exclusive
    context: str


@dataclass
class ScannedTemplate:
    fields: dict[str, dict] = field(default_factory=dict)
    occurrences: list[FieldOccurrence] = field(default_factory=list)
    preview_text: str = ""
    checksum: str = ""


def file_checksum(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(65536), b""):
            digest.update(chunk)
    return digest.hexdigest()


PLACEHOLDER_FIND = re.compile(
    r"\{\{[^{}]+\}\}|\[\[[^\[\]]+\]\]|\[[^\[\]]{1,80}\]"
)


def ensure_placeholder_highlights(path: Path) -> int:
    """Yellow-highlight [bracket] / {{mustache}} placeholders so they become fillable fields.

    Used when a user uploads a plain .docx as a template. Returns how many placeholders
    were highlighted (0 if the document already used yellow fields only).
    """
    document = Document(str(path))
    highlighted = 0

    for paragraph in _iter_paragraphs(document):
        text = paragraph.text or ""
        matches = list(PLACEHOLDER_FIND.finditer(text))
        if not matches:
            continue
        # Skip rebuild when every match is already covered by yellow runs
        already = scan_paragraph_yellow_placeholders(paragraph)
        if already and len(already) >= len(matches):
            continue
        # Preserve nothing fancy — rebuild runs with yellow on placeholders
        p_el = paragraph._p
        for child in list(p_el):
            if child.tag == qn("w:r"):
                p_el.remove(child)
        cursor = 0
        for match in matches:
            if match.start() > cursor:
                paragraph.add_run(text[cursor : match.start()])
            run = paragraph.add_run(match.group(0))
            run.font.highlight_color = YELLOW
            highlighted += 1
            cursor = match.end()
        if cursor < len(text):
            paragraph.add_run(text[cursor:])

    if highlighted:
        document.save(str(path))
    return highlighted


def scan_paragraph_yellow_placeholders(paragraph: Paragraph) -> list[str]:
    found: list[str] = []
    runs = paragraph.runs
    i = 0
    while i < len(runs):
        if not _is_yellow_run(runs[i]):
            i += 1
            continue
        chunks: list[str] = []
        while i < len(runs) and _is_yellow_run(runs[i]):
            chunks.append(runs[i].text or "")
            i += 1
        placeholder = "".join(chunks).strip()
        if placeholder:
            found.append(placeholder)
    return found


def scan_docx(path: Path) -> ScannedTemplate:
    document = Document(str(path))
    result = ScannedTemplate(checksum=file_checksum(path))
    preview_parts: list[str] = []

    for p_index, paragraph in enumerate(_iter_paragraphs(document)):
        text = paragraph.text.strip()
        if text:
            preview_parts.append(text)

        runs = paragraph.runs
        i = 0
        while i < len(runs):
            if not _is_yellow_run(runs[i]):
                i += 1
                continue
            start = i
            chunks: list[str] = []
            while i < len(runs) and _is_yellow_run(runs[i]):
                chunks.append(runs[i].text or "")
                i += 1
            placeholder = "".join(chunks)
            if not placeholder.strip():
                continue
            key, label = derive_field_key(placeholder)
            occurrence = FieldOccurrence(
                key=key,
                label=label,
                placeholder=placeholder.strip(),
                paragraph_index=p_index,
                run_start=start,
                run_end=i,
                context=text[:240],
            )
            result.occurrences.append(occurrence)
            bucket = result.fields.setdefault(
                key,
                {
                    "key": key,
                    "label": label,
                    "placeholder": placeholder.strip(),
                    "occurrence_count": 0,
                    "required": True,
                    "field_type": "short_text",
                    "sample_context": text[:240] or None,
                },
            )
            bucket["occurrence_count"] += 1

    result.preview_text = "\n\n".join(preview_parts[:40])
    return result


def fill_docx(
    source: Path,
    destination: Path,
    values: dict[str, str],
    *,
    clear_highlight: bool = True,
) -> list[str]:
    """Fill yellow fields by key. Returns list of still-missing required keys."""
    document = Document(str(source))
    filled_keys: set[str] = set()
    seen_keys: set[str] = set()

    for paragraph in _iter_paragraphs(document):
        runs = paragraph.runs
        i = 0
        while i < len(runs):
            if not _is_yellow_run(runs[i]):
                i += 1
                continue
            start = i
            chunks: list[str] = []
            while i < len(runs) and _is_yellow_run(runs[i]):
                chunks.append(runs[i].text or "")
                i += 1
            placeholder = "".join(chunks)
            if not placeholder.strip():
                continue
            key, _label = derive_field_key(placeholder)
            seen_keys.add(key)
            value = values.get(key)
            if value is None or value == "":
                continue
            # Put full replacement text on first run; clear the rest
            runs[start].text = value
            for run in runs[start + 1 : i]:
                run.text = ""
            if clear_highlight:
                for run in runs[start:i]:
                    run.font.highlight_color = None
            filled_keys.add(key)

    destination.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(destination))
    # python-docx strips theme/settings — restore a full package for EigenPal
    from app.services.docx_package import enrich_docx_for_editor

    enrich_docx_for_editor(destination)
    return sorted(seen_keys - filled_keys)
